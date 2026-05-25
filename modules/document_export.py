from pathlib import Path

from docx import Document


def export_to_docx(text, output_path):
    document = Document()
    document.add_heading("Распознанный рукописный текст", level=1)

    for paragraph in text.splitlines() or [""]:
        document.add_paragraph(paragraph)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path
